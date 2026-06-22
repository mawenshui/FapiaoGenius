"""Excel 导出服务"""

from pathlib import Path
from datetime import datetime
from typing import Optional

from core.logger import logger
from core.constants import TABLE_COLUMNS
from models.invoice import Invoice


class ExportService:
    """Excel 导出服务"""
    
    def export_to_excel(
        self,
        invoices: list[Invoice],
        output_path: str,
        title: Optional[str] = None
    ) -> bool:
        """
        导出发票数据到 Excel
        
        Args:
            invoices: 发票列表
            output_path: 输出文件路径
            title: 表格标题
            
        Returns:
            bool: 是否成功
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "发票数据"
            
            # 设置标题
            if title:
                ws.merge_cells('A1:M1')
                cell = ws['A1']
                cell.value = title
                cell.font = Font(size=14, bold=True)
                cell.alignment = Alignment(horizontal='center')
                start_row = 3
            else:
                start_row = 1
            
            # 写入表头
            headers = [col[1] for col in TABLE_COLUMNS if col[0] != 'id']
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_alignment = Alignment(horizontal='center', vertical='center')
            
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=start_row, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            
            # 写入数据
            for row_idx, invoice in enumerate(invoices, start_row + 1):
                data = [
                    invoice.business_type,
                    invoice.invoice_type,
                    invoice.invoice_number,
                    invoice.invoice_date,
                    invoice.buyer_name,
                    invoice.seller_name,
                    invoice.amount_without_tax,
                    invoice.tax_amount,
                    invoice.total_amount,
                    invoice.reimbursement_status,
                    invoice.status,
                    invoice.remark,
                ]
                
                for col_idx, value in enumerate(data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    
                    # 金额列格式化
                    if col_idx in [7, 8, 9]:  # 金额列
                        cell.number_format = '#,##0.00'
                        cell.alignment = Alignment(horizontal='right')
            
            # 设置列宽
            col_widths = [12, 12, 18, 12, 18, 18, 12, 10, 12, 10, 8, 18]
            for col_idx, width in enumerate(col_widths, 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width
            
            # 添加边框
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            for row in ws.iter_rows(min_row=start_row, max_row=start_row + len(invoices),
                                    min_col=1, max_col=len(headers)):
                for cell in row:
                    cell.border = thin_border
            
            # 保存文件
            wb.save(output_path)
            logger.info(f"Excel 导出成功: {output_path}")
            return True
            
        except ImportError:
            logger.error("openpyxl 未安装，无法导出 Excel")
            return False
        except Exception as e:
            logger.error(f"Excel 导出失败: {e}")
            return False
    
    def get_default_filename(self) -> str:
        """获取默认导出文件名"""
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        return f"发票数据_{timestamp}.xlsx"


# 全局导出服务实例
export_service = ExportService()
